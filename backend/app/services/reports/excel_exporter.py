import os
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from typing import Dict, Any

class ExcelExporter:
    @staticmethod
    def export(data: Dict[str, Any], output_path: str) -> str:
        os.makedirs(os.path.dirname(output_path), exist_ok=True)
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Executive Summary"

        header_font = Font(name="Calibri", size=14, bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="0284C7", end_color="0284C7", fill_type="solid")

        section_font = Font(name="Calibri", size=12, bold=True, color="0F172A")

        # Title
        ws.merge_cells("A1:E1")
        ws["A1"] = data.get("title", "Hospital Executive Report")
        ws["A1"].font = header_font
        ws["A1"].fill = header_fill
        ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[1].height = 30

        # Metadata
        ws["A3"] = "Report ID:"
        ws["B3"] = data.get("report_id")
        ws["A4"] = "Generated:"
        ws["B4"] = data.get("generated_at")
        ws["A5"] = "Type:"
        ws["B5"] = data.get("report_type")

        # Summary
        ws["A7"] = "Executive Summary"
        ws["A7"].font = section_font
        ws["A8"] = data.get("executive_summary", "")

        # Metrics
        ws["A10"] = "Key Performance Metrics"
        ws["A10"].font = section_font

        row = 11
        ws.cell(row=row, column=1, value="Metric").font = Font(bold=True)
        ws.cell(row=row, column=2, value="Value").font = Font(bold=True)
        row += 1

        metrics = data.get("metrics", {})
        for k, v in metrics.items():
            ws.cell(row=row, column=1, value=str(k).replace('_', ' ').title())
            ws.cell(row=row, column=2, value=str(v))
            row += 1

        # Data Sheet
        records = data.get("data", [])
        if records and isinstance(records, list):
            ws_data = wb.create_sheet(title="Data Records")
            keys = list(records[0].keys())

            for col_num, k in enumerate(keys, 1):
                cell = ws_data.cell(row=1, column=col_num, value=str(k).replace('_', ' ').title())
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill(start_color="334155", end_color="334155", fill_type="solid")

            for row_num, rec in enumerate(records, 2):
                for col_num, k in enumerate(keys, 1):
                    ws_data.cell(row=row_num, column=col_num, value=rec.get(k, ""))

        wb.save(output_path)
        return output_path
